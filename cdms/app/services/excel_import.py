"""Excel (.xlsx) parsing for the file-upload ingestion mechanism."""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook

HEADER_MAP: dict[str, str] = {
    "partnersku": "partnerSKU",
    "sku": "sku",
    "externalid": "external_id",
    "id": "external_id",
    "productname": "productName",
    "name": "productName",
    "color": "color",
    "size": "size",
    "description": "description",
    "assettype": "assetType",
    "isactive": "isActive",
    "price": "price",
    "stock": "stock",
    "quantity": "stock",
    "qty": "stock",
}

_TRUE = {"true", "1", "yes", "y", "active"}
_FALSE = {"false", "0", "no", "n", "inactive"}

ALLOWED_EXTENSIONS = (".xlsx",)
MAX_UPLOAD_BYTES = 5 * 1024 * 1024


@dataclass
class ParsedRow:
    row_number: int
    product: dict[str, Any] | None = None
    error: str | None = None


@dataclass
class ParsedWorkbook:
    rows: list[ParsedRow] = field(default_factory=list)

    @property
    def products(self) -> list[dict[str, Any]]:
        return [r.product for r in self.rows if r.product is not None]

    @property
    def errors(self) -> list[ParsedRow]:
        return [r for r in self.rows if r.error]


def normalize_header(raw: Any) -> str:
    return str(raw or "").strip().lower().replace(" ", "").replace("_", "").replace("-", "")


def _coerce(raw: Any, target_type: type | tuple[type, ...] | None = None) -> Any:
    if raw is None:
        return None
    if isinstance(raw, str):
        raw = raw.strip()
        if raw == "":
            return None
    types: tuple[type, ...] = (
        target_type if isinstance(target_type, tuple) else ((target_type,) if target_type else ())
    )
    if bool in types:
        if isinstance(raw, bool):
            return raw
        lowered = str(raw).strip().lower()
        if lowered in _TRUE:
            return True
        if lowered in _FALSE:
            return False
        return None
    if int in types or float in types:
        if isinstance(raw, bool):
            return None
        try:
            value = float(raw)
        except (TypeError, ValueError):
            return None
        if int in types and float not in types:
            return int(value) if value.is_integer() else None
        return int(value) if value.is_integer() else value
    return raw


_FIELD_TYPES: dict[str, type] = {
    "isActive": bool,
    "price": (int, float),
    "stock": int,
}


def parse_excel(data: bytes) -> ParsedWorkbook:
    if not data:
        raise ValueError("empty file")

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"invalid xlsx file: {exc}") from exc

    parsed = ParsedWorkbook()
    try:
        sheet = workbook.worksheets[0]
        row_iter = sheet.iter_rows(values_only=True)
        try:
            headers_raw = next(row_iter)
        except StopIteration:
            raise ValueError("worksheet is empty") from None

        columns: list[str | None] = []
        for h in headers_raw:
            normalized = normalize_header(h)
            columns.append(HEADER_MAP.get(normalized))

        if not any(columns):
            raise ValueError("no recognizable columns in header row")

        for row_number, values in enumerate(row_iter, start=2):
            if values is None or all(v is None or str(v).strip() == "" for v in values):
                continue
            product: dict[str, Any] = {}
            for column, value in zip(columns, values):
                if column is None:
                    continue
                coerced = _coerce(value, _FIELD_TYPES.get(column, str))
                if coerced is not None:
                    product[column] = coerced
            parsed.rows.append(ParsedRow(row_number=row_number, product=product))

        for row in parsed.rows:
            if row.product is None:
                continue
            has_id = any(
                str(row.product.get(k) or "").strip()
                for k in ("external_id", "partnerSKU", "sku")
            )
            if not has_id:
                row.error = "missing partnerSKU/sku"
                row.product = None
    finally:
        workbook.close()

    return parsed
